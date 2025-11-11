"""DOCX to markdown converter using python-docx.

Local DOCX extraction for air-gapped operation.
"""

import asyncio
from typing import Dict, List
import logging

from .base import BaseTextConverter

logger = logging.getLogger(__name__)


class DocxConverter(BaseTextConverter):
    """Converts DOCX files to markdown using python-docx (local extraction)."""

    async def convert_batch(self, file_paths: List[str]) -> Dict[str, str]:
        """Convert DOCX files to markdown text.

        Args:
            file_paths: List of DOCX file paths to convert

        Returns:
            Dict mapping file_path -> markdown text content (None if failed)
        """
        try:
            from docx import Document
        except ImportError:
            logger.error("python-docx package not installed for DOCX conversion")
            return {path: None for path in file_paths}

        logger.info(f"Converting {len(file_paths)} DOCX files to markdown...")

        results = {}
        semaphore = asyncio.Semaphore(10)  # Limit concurrent DOCX processing

        async def _convert_one(path: str):
            async with semaphore:
                try:
                    def _extract():
                        # Load DOCX document
                        doc = Document(path)

                        # Extract paragraphs
                        text_parts = []
                        for paragraph in doc.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                text_parts.append(text)

                        # Also extract text from tables
                        for table in doc.tables:
                            table_rows = []
                            for row in table.rows:
                                row_cells = [cell.text.strip() for cell in row.cells]
                                table_rows.append(row_cells)

                            # Convert table to markdown
                            if table_rows:
                                # Header row
                                header = table_rows[0]
                                text_parts.append("| " + " | ".join(header) + " |")
                                text_parts.append("| " + " | ".join(["---"] * len(header)) + " |")

                                # Data rows
                                for row in table_rows[1:]:
                                    padded = row + [""] * (len(header) - len(row))
                                    text_parts.append("| " + " | ".join(padded[:len(header)]) + " |")

                                text_parts.append("")  # Blank line after table

                        full_text = "\n\n".join(text_parts)
                        return full_text.strip() if full_text else None

                    text = await asyncio.to_thread(_extract)

                    if text:
                        results[path] = text
                        logger.debug(f"Converted DOCX file: {path} ({len(text)} characters)")
                    else:
                        logger.warning(f"DOCX conversion produced no content for {path}")
                        results[path] = None

                except Exception as e:
                    logger.error(f"DOCX conversion failed for {path}: {e}")
                    results[path] = None

        await asyncio.gather(*[_convert_one(p) for p in file_paths], return_exceptions=True)

        successful = sum(1 for r in results.values() if r)
        logger.info(f"DOCX conversion complete: {successful}/{len(file_paths)} files successful")

        return results
